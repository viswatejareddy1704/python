from openpyxl import Workbook,load_workbook

# wb=Workbook()
# ws=wb.active
# ws.title='students'
# wb.save('allstudentdetails.xlsx')
# wb.create_sheet('dashboard')
# wb.create_sheet('reports')
# wb.save('allstudentdeatils.xlsx')

# print(wb.sheetnames)

wb=load_workbook('allstudentdetails.xlsx')
# current=wb.active
# print(f"current working sheet name: {current.title}")
# print(f"all sheet names: {wb.sheetnames}")

ws=wb['dashboard']
ws=wb.sheetnames[1]
print(f"current working sheet name: {ws.title}")